import re
from pathlib import Path

from config.settings import ROOT_DIR, settings
from helpers.logger import get_logger

log = get_logger("exceldata")

HEADER_ROW = 4
_CACHE = {}


DEFAULT_FILE = "test_data/swing_ios_test_data.xlsx"


def workbook_path():
    path = Path(getattr(settings, "TEST_DATA_FILE", None) or DEFAULT_FILE)
    return path if path.is_absolute() else ROOT_DIR / path


def _key(header):
    cleaned = re.sub(r"[^a-zA-Z0-9]+", "_", str(header)).strip("_")
    return cleaned.upper()


def _read_rows(sheet):
    from openpyxl import load_workbook

    path = workbook_path()
    if not path.exists():
        raise FileNotFoundError(f"test data workbook not found: {path}")
    book = load_workbook(path, data_only=True, read_only=True)
    if sheet not in book.sheetnames:
        raise KeyError(f"sheet {sheet} not in {path.name}, sheets: {book.sheetnames}")
    grid = book[sheet]
    rows = []
    headers = []
    for index, values in enumerate(grid.iter_rows(values_only=True), start=1):
        if index < HEADER_ROW:
            continue
        if index == HEADER_ROW:
            headers = [_key(value) if value else "" for value in values]
            continue
        payload = {}
        for header, value in zip(headers, values):
            if not header or value is None or str(value).strip() == "":
                continue
            payload[header] = value.strip() if isinstance(value, str) else value
        if payload.get("TC_ID"):
            rows.append(payload)
    book.close()
    log.info(f"loaded {len(rows)} rows from {path.name}[{sheet}]")
    return rows


def _read_sheet(sheet):
    return {str(payload["TC_ID"]): payload for payload in _read_rows(sheet)}


def sheet_rows(sheet, refresh=False):
    if refresh or sheet not in _CACHE:
        _CACHE[sheet] = _read_sheet(sheet)
    return _CACHE[sheet]


def all_rows(sheet, refresh=False):
    key = f"{sheet}::rows"
    if refresh or key not in _CACHE:
        _CACHE[key] = _read_rows(sheet)
    return [dict(payload) for payload in _CACHE[key]]


def find_rows(sheet, column, value):
    wanted = str(value).strip()
    return [payload for payload in all_rows(sheet) if str(payload.get(column, "")).strip() == wanted]


def row(sheet, tc_id):
    rows = sheet_rows(sheet)
    if str(tc_id) not in rows:
        raise KeyError(f"{tc_id} not in {workbook_path().name}[{sheet}], ids: {sorted(rows)}")
    return dict(rows[str(tc_id)])


META = ("SHEET", "ALIASES", "META", "TC_ID")


class ExcelData:
    SHEET = None
    ALIASES = {}
    TC_ID = None
    _loaded = {}

    @classmethod
    def defaults(cls):
        if "_defaults" not in vars(cls):
            cls._defaults = {
                key: value
                for key, value in vars(cls).items()
                if key.isupper() and key not in META
            }
        return cls._defaults

    @classmethod
    def load(cls, tc_id):
        defaults = cls.defaults()
        for key in list(cls._loaded):
            if key in defaults:
                setattr(cls, key, defaults[key])
            elif key in vars(cls):
                delattr(cls, key)
        payload = row(cls.SHEET, tc_id)
        applied = {}
        for header, value in payload.items():
            name = cls.ALIASES.get(header, header)
            setattr(cls, name, value)
            applied[name] = value
        cls._loaded = applied
        cls.TC_ID = str(tc_id)
        return cls

    @classmethod
    def get(cls, name, default=None):
        return getattr(cls, name, default)

    @classmethod
    def as_dict(cls):
        return dict(cls._loaded)

    @classmethod
    def ids(cls):
        return sorted(sheet_rows(cls.SHEET))
